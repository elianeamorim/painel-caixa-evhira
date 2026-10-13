"""Gera os arquivos de amostra (fictícios) em testes/amostras/ e amostras_demo/.

Uso, na raiz do projeto:  .venv\\Scripts\\python scripts/gerar_amostras.py
O recorte é o "Banco A" da fábrica de exemplo: uma planilha de lançamentos no layout do Omie e o
extrato do mesmo banco em dois formatos (OFX e CSV), para subir arquivo de verdade na demo.
"""
import csv
import shutil
import sys
from datetime import datetime
from pathlib import Path

RAIZ = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(RAIZ))

from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Font  # noqa: E402

from core.exemplo import movimentos  # noqa: E402

# Layout imitado da exportação "Contas a Pagar" do Omie: título na linha 1, cabeçalho de 30 colunas na 2,
# totais na 3 (primeira coluna vazia), dados da 4 em diante, valores negativos. Só as 7 colunas lidas pelo
# sistema (Razão Social, Previsão, Vencimento, Valor Líquido, Conta Corrente, Documento, Categoria) vêm do
# arquivo real; o nome das outras é inventado.
COLUNAS_OMIE = [
    "Situação", "Razão Social", "Nome Fantasia", "CNPJ/CPF", "Documento", "Parcela", "Categoria",
    "Departamento", "Projeto", "Conta Corrente", "Emissão", "Entrada", "Vencimento", "Previsão",
    "Valor Original", "Juros", "Multa", "Desconto", "Valor Líquido", "Valor Pago", "Observação",
    "Tipo de Documento", "Número do Pedido", "Nota Fiscal", "Código de Barras", "Vendedor",
    "Registrado por", "Alterado por", "Forma de Pagamento", "Origem",
]
assert len(COLUNAS_OMIE) == 30


def escrever_omie(lancamentos, destino):
    wb = Workbook()
    ws = wb.active
    ws.title = "Contas a Pagar"
    pos = {nome: i for i, nome in enumerate(COLUNAS_OMIE)}
    ws.append(["Contas a Pagar gerado há 2 minutos"])
    ws.append(COLUNAS_OMIE)
    total = -sum(l["centavos"] for l in lancamentos) / 100
    totais = [None] * 30
    for nome in ("Valor Original", "Valor Líquido", "Valor Pago"):
        totais[pos[nome]] = total
    ws.append(totais)
    for n, l in enumerate(lancamentos):
        linha = [None] * 30
        linha[pos["Situação"]] = "Pago"
        linha[pos["Razão Social"]] = l["descricao"]
        linha[pos["Nome Fantasia"]] = l["descricao"].split(" ")[0]
        linha[pos["CNPJ/CPF"]] = f"11.{n % 1000:03d}.{(n * 7) % 1000:03d}/0001-00"
        linha[pos["Documento"]] = l["documento"]
        linha[pos["Parcela"]] = "1/1"
        linha[pos["Categoria"]] = l["categoria"]
        linha[pos["Conta Corrente"]] = l["conta"]
        linha[pos["Emissão"]] = datetime(l["data"].year, l["data"].month, 1)
        linha[pos["Vencimento"]] = datetime.combine(l["data"], datetime.min.time())
        linha[pos["Previsão"]] = datetime.combine(l["data"], datetime.min.time())
        for nome in ("Valor Original", "Valor Líquido", "Valor Pago"):
            linha[pos[nome]] = -l["centavos"] / 100
        linha[pos["Forma de Pagamento"]] = "Boleto"
        ws.append(linha)
    for celula in ws[2]:
        celula.font = Font(bold=True)
    for col in ("K", "L", "M", "N"):
        for c in ws[col][3:]:
            c.number_format = "dd/mm/yyyy"
    ws.freeze_panes = "A4"
    wb.save(destino)


def escrever_ofx(extrato, destino):
    linhas = [
        "OFXHEADER:100", "DATA:OFXSGML", "VERSION:102", "SECURITY:NONE", "ENCODING:USASCII",
        "CHARSET:1252", "COMPRESSION:NONE", "OLDFILEUID:NONE", "NEWFILEUID:NONE", "",
        "<OFX>", "<SIGNONMSGSRSV1><SONRS><STATUS><CODE>0<SEVERITY>INFO</STATUS>"
        "<DTSERVER>20261001000000<LANGUAGE>POR</SONRS></SIGNONMSGSRSV1>",
        "<BANKMSGSRSV1><STMTTRNRS><TRNUID>1<STATUS><CODE>0<SEVERITY>INFO</STATUS>",
        "<STMTRS><CURDEF>BRL<BANKACCTFROM><BANKID>0000<ACCTID>00000-0<ACCTTYPE>CHECKING</BANKACCTFROM>",
        "<BANKTRANLIST><DTSTART>20260901<DTEND>20260930",
    ]
    for n, e in enumerate(extrato):
        tipo = "CREDIT" if e["centavos"] > 0 else "DEBIT"
        linhas += [
            "<STMTTRN>", f"<TRNTYPE>{tipo}", f"<DTPOSTED>{e['data'].strftime('%Y%m%d')}120000[-3:BRT]",
            f"<TRNAMT>{e['centavos'] / 100:.2f}", f"<FITID>{20260900 + n}", f"<MEMO>{e['descricao']}",
            "</STMTTRN>",
        ]
    linhas += ["</BANKTRANLIST></STMTRS></STMTTRNRS></BANKMSGSRSV1>", "</OFX>", ""]
    Path(destino).write_text("\r\n".join(linhas), encoding="cp1252")


def escrever_csv(extrato, destino):
    with open(destino, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f, delimiter=";")
        w.writerow(["Data", "Histórico", "Documento", "Valor"])
        for n, e in enumerate(extrato):
            valor = f"{e['centavos'] / 100:,.2f}".replace(",", "x").replace(".", ",").replace("x", ".")
            w.writerow([e["data"].strftime("%d/%m/%Y"), e["descricao"], 700000 + n, valor])


def main():
    lancamentos, extrato = movimentos("Banco A")
    pasta = RAIZ / "testes" / "amostras"
    pasta.mkdir(parents=True, exist_ok=True)
    escrever_omie(lancamentos, pasta / "lancamentos_omie_exemplo.xlsx")
    escrever_ofx(extrato, pasta / "extrato_bancoA_exemplo.ofx")
    escrever_csv(extrato, pasta / "extrato_bancoA_exemplo.csv")
    demo = RAIZ / "amostras_demo"
    demo.mkdir(exist_ok=True)
    for arquivo in pasta.iterdir():
        shutil.copy2(arquivo, demo / arquivo.name)
    print(f"{len(lancamentos)} lançamentos e {len(extrato)} movimentos do extrato gravados em {pasta} e {demo}")


if __name__ == "__main__":
    main()
