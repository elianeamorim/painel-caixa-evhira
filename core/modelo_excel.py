"""Planilha modelo dos lançamentos: cabeçalho na linha 1, exemplo na linha 2, validações do Excel.

Gerada na hora (não é arquivo estático). O leitor reconhece este cabeçalho como "Planilha modelo Évhira".
"""
import io
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

COR_CABECALHO = "866E29"  # mesmo valor de --accent-texto em static/css/tokens.css (a planilha não lê CSS)
COLUNAS = [("Data", 14), ("Descrição", 42), ("Valor", 16), ("Conta", 18), ("Categoria", 22), ("Documento", 16)]
EXEMPLO = [date(2026, 9, 5), "Fornecedor Exemplo Ltda", 1234.56, "Banco A", "Custos de produção", "NF-1001"]
LINHAS_VALIDADAS = 2000
INSTRUCOES = [
    "Como preencher",
    "1. Uma linha por pagamento. A linha 2 é só um exemplo: apague ou troque.",
    "2. Data: o dia em que o dinheiro saiu (ou vai sair), no formato dd/mm/aaaa.",
    "3. Valor: quanto foi pago, em reais. Pode ser com ou sem sinal de menos.",
    "4. Conta: o banco de onde saiu (Banco A, Banco B...). Pagamento no cartão, na caixinha ou adiantamento: escreva isso aqui.",
    "5. Categoria e Documento são opcionais, mas ajudam a achar o lançamento depois.",
    "6. Não mude os títulos da linha 1 nem a ordem das colunas.",
]


def gerar_modelo():
    """Devolve os bytes do .xlsx."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Lançamentos"
    ws.append([nome for nome, _ in COLUNAS])
    ws.append(EXEMPLO)
    for i, (_, largura) in enumerate(COLUNAS, start=1):
        celula = ws.cell(row=1, column=i)
        celula.font = Font(bold=True, color="FFFFFF")
        celula.fill = PatternFill("solid", fgColor=COR_CABECALHO)
        celula.alignment = Alignment(vertical="center")
        ws.column_dimensions[celula.column_letter].width = largura
    ws["A2"].number_format = "dd/mm/yyyy"
    ws["C2"].number_format = "#,##0.00"
    ws.freeze_panes = "A2"

    fim = LINHAS_VALIDADAS
    validacoes = [
        (DataValidation(type="date", operator="greaterThanOrEqual", formula1="36526", allow_blank=True,
                        errorTitle="Data inválida", error="Digite uma data, por exemplo 05/09/2026."), f"A2:A{fim}"),
        (DataValidation(type="textLength", operator="lessThanOrEqual", formula1="120", allow_blank=True,
                        errorTitle="Texto longo demais", error="Use no máximo 120 letras."), f"B2:B{fim}"),
        (DataValidation(type="decimal", operator="between", formula1="-1000000000", formula2="1000000000",
                        allow_blank=True, errorTitle="Valor inválido",
                        error="Digite só números, por exemplo 1234,56."), f"C2:C{fim}"),
    ]
    for validacao, faixa in validacoes:
        validacao.showErrorMessage = True
        ws.add_data_validation(validacao)
        validacao.add(faixa)

    ajuda = wb.create_sheet("Como preencher")
    for linha in INSTRUCOES:
        ajuda.append([linha])
    ajuda["A1"].font = Font(bold=True, size=13)
    ajuda.column_dimensions["A"].width = 110

    saida = io.BytesIO()
    wb.save(saida)
    return saida.getvalue()
