"""Leitura dos arquivos da conferência.

Tudo acontece em memória: o arquivo chega como bytes, vira uma lista de lançamentos e é
descartado. Nada é gravado em disco e nenhum conteúdo de linha vai pra log.
O formato é reconhecido pelo CONTEÚDO do arquivo, nunca pela extensão.
"""
import csv
import io
import re
import unicodedata
import zipfile
from datetime import date, datetime
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from itertools import islice

from openpyxl import load_workbook

MAX_LINHAS = 5000                      # decisão 2A: tabela roda no navegador, volume de centenas de linhas
MAX_DESCOMPRIMIDO = 50 * 1024 * 1024   # barreira contra xlsx "bomba" (zip pequeno que abre gigante)
LINHAS_DE_CABECALHO = 10               # o cabeçalho é procurado nas 10 primeiras linhas

NOMES_FORMATO = {
    "ofx": "OFX do banco",
    "omie": "Omie (contas a pagar)",
    "modelo": "Planilha modelo Évhira",
    "planilha": "Planilha com colunas de data, descrição e valor",
}

MSG_SEM_COLUNAS = (
    "Não encontrei as colunas de data, descrição e valor neste arquivo. "
    "Use a exportação do seu sistema ou baixe o modelo e preencha."
)
MSG_NAO_LIDO = "Não consegui abrir este arquivo. Confira se ele não está corrompido ou baixe o modelo."

# Sinônimos de cabeçalho, em ordem de preferência (já sem acento, em minúsculas).
SINONIMOS = {
    "data": ["previsao", "data pagamento", "data de pagamento", "data do pagamento", "pagamento",
             "vencimento", "data vencimento", "data de vencimento", "data lancamento", "data movimento",
             "data"],
    "descricao": ["razao social", "descricao", "historico", "descricao do lancamento", "lancamento",
                  "favorecido", "fornecedor", "beneficiario", "nome", "memo", "detalhes"],
    "valor": ["valor liquido", "valor pago", "valor", "montante"],
    "conta": ["conta corrente", "conta", "banco"],
    "documento": ["documento", "numero documento", "nota fiscal"],
    "categoria": ["categoria", "grupo", "classificacao"],
}
CABECALHO_MODELO = {"data", "descricao", "valor", "conta", "categoria", "documento"}
FORA_DO_BANCO = ("cartao", "caixinha", "adiantamento", "dinheiro")


class ErroArquivo(Exception):
    """Erro de domínio: a mensagem foi escrita pra aparecer na tela do dono, sem jargão."""


def normalizar(texto):
    """Minúsculas, sem acento, só letras e números separados por espaço."""
    t = unicodedata.normalize("NFD", str(texto if texto is not None else ""))
    t = "".join(c for c in t if not unicodedata.combining(c))
    return re.sub(r"[^a-z0-9]+", " ", t.lower()).strip()


def passa_pelo_banco(conta):
    """Lançamento em cartão, caixinha ou adiantamento não aparece no extrato como linha própria."""
    # escada: lista fixa de palavras; virar configuração se outro sistema nomear a conta de outro jeito
    c = normalizar(conta)
    if not c:
        return True
    return not (c == "caixa" or any(p in c for p in FORA_DO_BANCO))


def parse_data(valor):
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    if not isinstance(valor, str):
        return None
    texto = valor.strip().replace("T", " ").split(" ")[0]
    for formato in ("%d/%m/%Y", "%d/%m/%y", "%Y-%m-%d", "%d-%m-%Y", "%d.%m.%Y", "%Y%m%d"):
        try:
            d = datetime.strptime(texto, formato).date()
        except ValueError:
            continue
        return d if 1990 <= d.year <= 2100 else None
    return None


def parse_centavos(valor):
    """'1.234,56', '-1234.56', '(1.234,56)', 'R$ 12,00' e números viram centavos inteiros com sinal."""
    if valor is None or isinstance(valor, bool):
        return None
    if isinstance(valor, (int, float)):
        numero = Decimal(str(valor))
        negativo = False
    else:
        s = str(valor).strip()
        negativo = s.startswith("-") or s.endswith("-") or (s.startswith("(") and s.endswith(")"))
        s = re.sub(r"[^\d.,]", "", s)
        if "," in s and "." in s:
            if s.rfind(",") > s.rfind("."):
                s = s.replace(".", "").replace(",", ".")
            else:
                s = s.replace(",", "")
        elif "," in s:
            s = s.replace(".", "").replace(",", ".")
        elif re.fullmatch(r"\d{1,3}(\.\d{3})+", s):
            s = s.replace(".", "")
        if not re.fullmatch(r"\d+(\.\d+)?", s):
            return None
        numero = Decimal(s)
    try:
        centavos = int((numero * 100).to_integral_value(ROUND_HALF_UP))
    except (InvalidOperation, ValueError):
        return None
    return -centavos if negativo else centavos


def _texto(valor, limite=200):
    return re.sub(r"\s+", " ", str(valor)).strip()[:limite] if valor is not None else ""


def _item(data, descricao, centavos, conta="", documento="", categoria=""):
    return {"data": data, "descricao": descricao or "Sem descrição", "centavos": centavos,
            "conta": conta, "documento": documento, "categoria": categoria}


# ---------- OFX ----------

def _ler_ofx(dados):
    try:
        texto = dados.decode("utf-8")
    except UnicodeDecodeError:
        texto = dados.decode("cp1252", errors="replace")
    itens = []
    for bloco in re.findall(r"<STMTTRN>(.*?)(?=</STMTTRN>|<STMTTRN>|</BANKTRANLIST>|$)", texto, re.S | re.I):
        def campo(tag):
            m = re.search(rf"<{tag}>([^<\r\n]*)", bloco, re.I)
            return m.group(1).strip() if m else ""
        data = parse_data(campo("DTPOSTED")[:8])
        centavos = parse_centavos(campo("TRNAMT"))
        if data is None or centavos is None:
            continue
        partes = []
        for p in (_texto(campo("NAME")), _texto(campo("MEMO"))):
            if p and p not in partes:
                partes.append(p)
        itens.append(_item(data, " ".join(partes), centavos))
    if not itens:
        raise ErroArquivo("Este arquivo OFX não trouxe nenhuma movimentação.")
    return itens


# ---------- planilhas (xlsx e csv) ----------

def _indices(cabecalho, candidatos):
    """Colunas do cabeçalho que casam com os sinônimos, na ordem de preferência."""
    achados = []
    for cand in candidatos:
        for i, h in enumerate(cabecalho):
            if i not in achados and (h == cand or (cand != "data" and h.startswith(cand + " "))):
                achados.append(i)
    return achados


def _ler_tabela(linhas):
    """Procura o cabeçalho nas primeiras linhas e lê o resto. None se não achar data, descrição e valor."""
    for n, linha in enumerate(linhas[:LINHAS_DE_CABECALHO]):
        cab = [normalizar(c) for c in linha]
        col = {campo: _indices(cab, SINONIMOS[campo]) for campo in SINONIMOS}
        if col["data"] and col["descricao"] and col["valor"]:
            break
    else:
        return None

    if "razao social" in cab and "conta corrente" in cab:
        formato = "omie"
    elif {c for c in cab if c} == CABECALHO_MODELO:
        formato = "modelo"
    else:
        formato = "planilha"

    def celula(linha, indices):
        return linha[indices[0]] if indices and indices[0] < len(linha) else None

    itens, ignoradas = [], 0
    for linha in linhas[n + 1:]:
        data = next((d for d in (parse_data(linha[i]) for i in col["data"] if i < len(linha)) if d), None)
        descricao = _texto(celula(linha, col["descricao"]))
        centavos = parse_centavos(celula(linha, col["valor"]))
        if data is None and not descricao:
            continue  # linha vazia ou de totais
        if data is None or centavos is None:
            ignoradas += 1
            continue
        itens.append(_item(data, descricao, centavos,
                           _texto(celula(linha, col["conta"])),
                           _texto(celula(linha, col["documento"])),
                           _texto(celula(linha, col["categoria"]))))
    if not itens:
        raise ErroArquivo("Achei as colunas, mas nenhuma linha com data e valor válidos.")
    if len(itens) > MAX_LINHAS:
        raise ErroArquivo(f"O arquivo tem mais de {MAX_LINHAS} linhas. Envie um período menor.")
    return formato, itens, ignoradas


def _folhas_xlsx(dados):
    try:
        with zipfile.ZipFile(io.BytesIO(dados)) as z:
            if sum(i.file_size for i in z.infolist()) > MAX_DESCOMPRIMIDO:
                raise ErroArquivo("Este arquivo é grande demais depois de aberto.")
        wb = load_workbook(io.BytesIO(dados), read_only=True, data_only=True)
        try:
            for ws in wb.worksheets:
                yield list(islice(ws.iter_rows(values_only=True), MAX_LINHAS + 20))
        finally:
            wb.close()
    except ErroArquivo:
        raise
    except Exception:  # fronteira de confiança: arquivo do usuário pode quebrar de muitas formas
        raise ErroArquivo(MSG_NAO_LIDO) from None


def _folhas_csv(dados):
    try:
        texto = dados.decode("utf-8-sig")
    except UnicodeDecodeError:
        texto = dados.decode("cp1252", errors="replace")
    for separador in (";", ",", "\t"):
        try:
            yield list(islice(csv.reader(io.StringIO(texto), delimiter=separador), MAX_LINHAS + 20))
        except csv.Error:
            raise ErroArquivo(MSG_NAO_LIDO) from None


def _detectar(dados):
    if dados[:4] == b"PK\x03\x04":
        return "xlsx"
    if dados[:4] == b"%PDF":
        raise ErroArquivo("PDF não dá pra ler aqui. Exporte o arquivo em planilha (xlsx ou csv) ou em OFX.")
    if dados[:8] == b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1":
        raise ErroArquivo("Esta planilha está no formato antigo (.xls). Abra no Excel e salve como .xlsx.")
    if b"\x00" in dados[:4096]:
        raise ErroArquivo(MSG_NAO_LIDO)
    cabeca = dados[:2048].upper()
    if b"OFXHEADER" in cabeca or b"<OFX>" in cabeca:
        return "ofx"
    return "texto"


def _ler_bruto(dados):
    """-> (formato, itens com sinal, linhas ignoradas)."""
    if not dados or not dados.strip():
        raise ErroArquivo("O arquivo está vazio.")
    tipo = _detectar(dados)
    if tipo == "ofx":
        return "ofx", _ler_ofx(dados), 0
    folhas = _folhas_xlsx(dados) if tipo == "xlsx" else _folhas_csv(dados)
    for linhas in folhas:
        achado = _ler_tabela(linhas)
        if achado:
            return achado
    raise ErroArquivo(MSG_SEM_COLUNAS)


# ---------- entradas públicas ----------

def ler_lancamentos(dados):
    """Lançamentos da contabilidade. Valor sempre em módulo; conta fora do banco vai pra 'nao_passam'."""
    formato, itens, ignoradas = _ler_bruto(dados)
    if formato == "ofx":
        raise ErroArquivo('Este arquivo parece o extrato do banco. Coloque ele no campo "Extrato do banco".')
    passam, fora = [], []
    for it in itens:
        it = {**it, "centavos": abs(it["centavos"])}
        (passam if passa_pelo_banco(it["conta"]) else fora).append(it)
    return {"formato": formato, "formato_nome": NOMES_FORMATO[formato],
            "itens": passam, "nao_passam": fora, "ignoradas": ignoradas}


def ler_extrato(dados):
    """Extrato do banco. Só as saídas (valores negativos) entram na conferência."""
    formato, itens, ignoradas = _ler_bruto(dados)
    if formato == "omie":
        raise ErroArquivo("Este arquivo parece a exportação de contas do seu sistema. "
                          'Coloque ele no campo "Lançamentos da contabilidade".')
    # escada: saída = valor negativo; extrato com Débito e Crédito em colunas separadas não entra, a tela de mapeamento (opção A) resolveria
    saidas = [{**i, "centavos": -i["centavos"]} for i in itens if i["centavos"] < 0]
    if not saidas:
        raise ErroArquivo("Não achei saídas neste extrato. As saídas precisam aparecer como valores negativos.")
    return {"formato": formato, "formato_nome": NOMES_FORMATO[formato], "itens": saidas,
            "fora_das_saidas": len(itens) - len(saidas), "ignoradas": ignoradas}
