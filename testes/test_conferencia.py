"""Testes do motor da conferência. Rodar na raiz do projeto: python testes/test_conferencia.py

Sem framework: cada função test_* faz asserts e o final conta quantas passaram.
"""
import io
import re
import sys
from datetime import date, datetime
from pathlib import Path

RAIZ = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(RAIZ))

from openpyxl import Workbook  # noqa: E402

from core.conferencia import conferir, parecida  # noqa: E402
from core.exemplo import BASE, dados_exemplo  # noqa: E402
from core.leitores import (ErroArquivo, ler_extrato, ler_lancamentos, parse_centavos,  # noqa: E402
                           parse_data, passa_pelo_banco)
from core.modelo_excel import gerar_modelo  # noqa: E402

AMOSTRAS = RAIZ / "testes" / "amostras"


def xlsx(linhas, folha="Dados"):
    wb = Workbook()
    ws = wb.active
    ws.title = folha
    for linha in linhas:
        ws.append(linha)
    saida = io.BytesIO()
    wb.save(saida)
    return saida.getvalue()


def lanc(dia, descricao, reais, conta=""):
    return {"data": date(2026, 9, dia), "descricao": descricao, "centavos": round(reais * 100), "conta": conta,
            "documento": "", "categoria": ""}


def pacote(*itens, nao_passam=()):
    return {"itens": list(itens), "nao_passam": list(nao_passam)}


def levanta(funcao, dados):
    try:
        funcao(dados)
    except ErroArquivo as e:
        assert str(e) and "Traceback" not in str(e) and "File " not in str(e)
        return str(e)
    raise AssertionError("era pra levantar ErroArquivo")


# ---------- números e datas ----------

def test_valores():
    casos = {"1.234,56": 123456, "-1234.56": -123456, "(1.234,56)": -123456, "R$ 12,00": 1200, "1234,5": 123450,
             "1.234": 123400, "0,99": 99, "1,234.56": 123456, "12": 1200, 12.5: 1250, -3.1: -310, "10,00-": -1000}
    for entrada, esperado in casos.items():
        assert parse_centavos(entrada) == esperado, (entrada, parse_centavos(entrada))
    for ruim in ("", "abc", None, True, "1,2,3x"):
        assert parse_centavos(ruim) is None, ruim


def test_datas():
    assert parse_data("03/09/2026") == date(2026, 9, 3)
    assert parse_data("03/09/26") == date(2026, 9, 3)
    assert parse_data("2026-09-03T10:00:00") == date(2026, 9, 3)
    assert parse_data("20260903120000[-3:BRT]"[:8]) == date(2026, 9, 3)
    assert parse_data(datetime(2026, 9, 3, 8, 0)) == date(2026, 9, 3)
    assert parse_data("31/02/2026") is None and parse_data("texto") is None and parse_data(None) is None


def test_contas_fora_do_banco():
    for conta in ("Cartão de Crédito", "Caixinha", "Adiantamento", "Caixa", "cartao visa"):
        assert not passa_pelo_banco(conta), conta
    for conta in ("Banco A", "Caixa Econômica", "", None):
        assert passa_pelo_banco(conta), conta


# ---------- leitores ----------

def test_ofx():
    texto = ("OFXHEADER:100\r\nDATA:OFXSGML\r\nCHARSET:1252\r\n\r\n<OFX><BANKMSGSRSV1><STMTTRNRS><STMTRS>"
             "<BANKTRANLIST>\r\n<STMTTRN><TRNTYPE>DEBIT<DTPOSTED>20260903120000[-3:BRT]<TRNAMT>-1234,56"
             "<FITID>1<MEMO>PAGTO BOLETO FORNECEDOR\r\n</STMTTRN>\r\n<STMTTRN><TRNTYPE>CREDIT<DTPOSTED>20260904"
             "<TRNAMT>500.00<FITID>2<NAME>PIX RECEBIDO\r\n</STMTTRN></BANKTRANLIST></STMTRS></STMTTRNRS>"
             "</BANKMSGSRSV1></OFX>")
    r = ler_extrato(texto.encode("cp1252"))
    assert r["formato"] == "ofx" and len(r["itens"]) == 1 and r["fora_das_saidas"] == 1
    assert r["itens"][0]["centavos"] == 123456 and r["itens"][0]["data"] == date(2026, 9, 3)
    assert "FORNECEDOR" in r["itens"][0]["descricao"]
    assert "extrato do banco" in levanta(ler_lancamentos, texto.encode("cp1252"))


def test_csv_com_ponto_e_virgula_e_acento():
    texto = "Data;Histórico;Valor\n03/09/2026;Pagamento Padaria Açúcar;-1.234,56\n04/09/2026;Pix recebido;500,00\n"
    for codificado in (texto.encode("utf-8"), texto.encode("cp1252"), b"\xef\xbb\xbf" + texto.encode("utf-8")):
        r = ler_extrato(codificado)
        assert r["formato"] == "planilha" and len(r["itens"]) == 1 and r["itens"][0]["centavos"] == 123456
        assert r["itens"][0]["descricao"].startswith("Pagamento Padaria")


def test_csv_com_virgula_e_cabecalho_depois_de_titulo():
    texto = 'Extrato da conta,,\nPeríodo setembro,,\nData,Descrição,Valor\n03/09/2026,"Posto, Estrada",-80.50\n'
    r = ler_extrato(texto.encode("utf-8"))
    assert r["itens"][0]["descricao"] == "Posto, Estrada" and r["itens"][0]["centavos"] == 8050


def test_xlsx_estilo_omie():
    # título na linha 1, cabeçalho de 30 colunas na 2, totais na 3 (primeira coluna vazia), dados da 4, valores negativos
    cab = ["Situação", "Razão Social", "Documento", "Categoria", "Conta Corrente", "Vencimento", "Previsão",
           "Valor Original", "Valor Líquido"] + [f"Coluna {i}" for i in range(21)]
    assert len(cab) == 30
    linhas = [["Contas a Pagar gerado há 3 minutos"], cab, [None] * 7 + [-900.0, -900.0] + [None] * 21]

    def dado(razao, venc, prev, valor, conta):
        return ["Pago", razao, "NF-1", "Frete", conta, venc, prev, valor, valor] + [None] * 21
    linhas += [
        dado("Transportes Alfa", datetime(2026, 9, 2), datetime(2026, 9, 4), -400.0, "Banco A"),
        dado("Posto Beta", datetime(2026, 9, 5), None, -300.0, "Cartão de Crédito"),   # Previsão vazia: usa Vencimento
        dado("Gráfica Gama", "07/09/2026", "08/09/2026", "-200,00", "Banco B"),
        ["Pago", "Linha sem data", None, None, "Banco A", None, None, "x", "abc"] + [None] * 21,  # ignorada, conta 1
    ]
    r = ler_lancamentos(xlsx(linhas))
    assert r["formato"] == "omie" and len(r["itens"]) == 2 and len(r["nao_passam"]) == 1 and r["ignoradas"] == 1
    por_nome = {i["descricao"]: i for i in r["itens"] + r["nao_passam"]}
    assert por_nome["Transportes Alfa"]["data"] == date(2026, 9, 4)       # Previsão vence Vencimento
    assert por_nome["Posto Beta"]["data"] == date(2026, 9, 5)             # Vencimento como reserva
    assert por_nome["Gráfica Gama"]["centavos"] == 20000                  # valor em módulo
    assert por_nome["Transportes Alfa"]["documento"] == "NF-1" and por_nome["Transportes Alfa"]["categoria"] == "Frete"
    assert "contas do seu sistema" in levanta(ler_extrato, xlsx(linhas))


def test_planilha_generica_com_sinonimos_na_segunda_folha():
    vazia = [["Nada aqui"]]
    util = [["Relatório"], ["Vencimento", "Histórico", "Valor pago"], [date(2026, 9, 3), "Aluguel", "1.500,00"]]
    wb = Workbook()
    wb.active.title = "Capa"
    wb.active.append(vazia[0])
    ws = wb.create_sheet("Contas")
    for linha in util:
        ws.append(linha)
    saida = io.BytesIO()
    wb.save(saida)
    r = ler_lancamentos(saida.getvalue())
    assert r["formato"] == "planilha" and r["itens"][0]["centavos"] == 150000 and r["itens"][0]["descricao"] == "Aluguel"


def test_modelo_excel_e_lido_de_volta():
    modelo = gerar_modelo()
    r = ler_lancamentos(modelo)
    assert r["formato"] == "modelo" and len(r["itens"]) == 1 and r["itens"][0]["centavos"] == 123456
    from openpyxl import load_workbook
    ws = load_workbook(io.BytesIO(modelo)).worksheets[0]
    assert ws.freeze_panes == "A2" and ws["A1"].value == "Data" and ws["B2"].value == "Fornecedor Exemplo Ltda"
    assert len(ws.data_validations.dataValidation) >= 2


def test_arquivo_nao_reconhecido_da_erro_claro():
    levanta(ler_lancamentos, b"")
    levanta(ler_lancamentos, b"   \n ")
    assert "PDF" in levanta(ler_lancamentos, b"%PDF-1.7 ...")
    assert ".xlsx" in levanta(ler_lancamentos, b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"0" * 20)
    levanta(ler_lancamentos, bytes(range(256)) * 4)                       # binário qualquer
    levanta(ler_lancamentos, b"PK\x03\x04 isto nao e um zip de verdade")   # assinatura de xlsx, conteúdo quebrado
    levanta(ler_lancamentos, "só texto solto\noutra linha\n".encode())
    levanta(ler_lancamentos, "Nome;Telefone\nAna;123\n".encode())         # planilha sem data e valor
    levanta(ler_lancamentos, "Data;Histórico;Valor\nsem data;x;sem valor\n".encode())
    levanta(ler_extrato, "Data;Histórico;Valor\n03/09/2026;Entrada;100,00\n".encode())  # sem saídas
    xlsx_ruim = xlsx([["a", "b"]])
    levanta(ler_extrato, xlsx_ruim)


def test_limite_de_linhas():
    linhas = [["Data", "Descrição", "Valor"]] + [[date(2026, 9, 1), f"Item {i}", -1.0] for i in range(5100)]
    assert "5000" in levanta(ler_lancamentos, xlsx(linhas))


# ---------- conferência ----------

def test_tres_grupos():
    lancamentos = pacote(
        lanc(3, "Fornecedor Azul Ltda", 100.00),
        lanc(5, "Padaria Doce Pão", 250.00),
        lanc(9, "Gráfica Cores", 80.00),
        nao_passam=[lanc(4, "Posto", 30.0, "Cartão de Crédito")])
    extrato = pacote(
        lanc(4, "PAGTO BOLETO FORNECEDOR AZUL", 100.00),
        lanc(6, "PIX ENVIADO PADARIA DOCE PAO", 205.00),
        lanc(15, "TARIFA", 9.90))
    r = conferir(lancamentos, extrato)
    assert [x["descricao"] for x in r["bateu"]] == ["Fornecedor Azul Ltda"]
    assert [(x["descricao"], x["diferenca"]) for x in r["valor_diferente"]] == [("Padaria Doce Pão", -45.0)]
    assert sorted((x["lado"], x["descricao"]) for x in r["so_um_lado"]) == [("banco", "TARIFA"), ("lancamento", "Gráfica Cores")]
    assert r["nao_passam"]["quantidade"] == 1 and r["nao_passam"]["total"] == 30.0
    assert r["resumo"]["lancamentos"] == 4 and r["resumo"]["precisam_de_atencao"] == 3


def test_janela_de_tres_dias():
    for dif, deve_bater in ((3, True), (-3, True), (4, False)):
        r = conferir(pacote(lanc(10, "Fornecedor Azul", 100.0)), pacote(lanc(10 + dif, "Fornecedor Azul", 100.0)))
        assert bool(r["bateu"]) == deve_bater, dif
        if not deve_bater:
            assert len(r["so_um_lado"]) == 2 and not r["valor_diferente"]


def test_um_pra_um_e_duplicado():
    r = conferir(pacote(lanc(3, "Embalagens Alfa", 50.0), lanc(3, "Embalagens Alfa", 50.0)),
                 pacote(lanc(3, "PIX ENVIADO EMBALAGENS ALFA", 50.0)))
    assert len(r["bateu"]) == 1 and len(r["so_um_lado"]) == 1
    assert r["so_um_lado"][0]["motivo"] == "Parece lançado em duplicidade"


def test_regra_de_descricao_parecida():
    assert parecida("Laticínios Vale Verde", "PAGTO BOLETO LATICINIOS VALE VERD")   # acento, caixa, truncado
    assert parecida("Aluguel", "PIX ENVIADO ALUGUEL GALPAO 123")                    # uma contém a outra, sem número
    assert parecida("Moinho Dourado", "TED MOINHO PEDRA")                           # palavra de 4+ letras em comum
    assert not parecida("Moinho Pedra", "PAGTO BOLETO TARIFA")                      # sem nada em comum
    assert not parecida("Pagamento fornecedor", "PAGAMENTO BOLETO")                 # só palavras comuns
    assert not parecida("", "ALGO") and not parecida("Ltda", "LTDA")


def test_valor_diferente_exige_descricao_parecida():
    r = conferir(pacote(lanc(3, "Fornecedor Azul", 100.0)), pacote(lanc(3, "TARIFA BANCARIA", 90.0)))
    assert not r["valor_diferente"] and len(r["so_um_lado"]) == 2


# ---------- dados de exemplo e amostras ----------

def test_exemplo_tem_os_erros_plantados():
    lancamentos, extrato = dados_exemplo()
    r = conferir(lancamentos, extrato)
    assert len(BASE) == 54 and r["resumo"]["lancamentos"] == 59
    assert len(r["bateu"]) == 52
    assert len(r["valor_diferente"]) == 2 and all(x["diferenca"] == -90.0 for x in r["valor_diferente"])
    assert [x["motivo"] for x in r["so_um_lado"] if x["lado"] == "lancamento"] == ["Parece lançado em duplicidade"] * 2
    assert sum(1 for x in r["so_um_lado"] if x["lado"] == "banco") == 3
    assert r["nao_passam"]["quantidade"] == 3
    assert {i["conta"] for i in lancamentos["itens"]} == {"Banco A", "Banco B"}


def test_amostras_em_arquivo():
    lancamentos = ler_lancamentos((AMOSTRAS / "lancamentos_omie_exemplo.xlsx").read_bytes())
    assert lancamentos["formato"] == "omie" and len(lancamentos["nao_passam"]) == 3
    for nome, formato in (("extrato_bancoA_exemplo.ofx", "ofx"), ("extrato_bancoA_exemplo.csv", "planilha")):
        extrato = ler_extrato((AMOSTRAS / nome).read_bytes())
        assert extrato["formato"] == formato
        r = conferir(lancamentos, extrato)
        assert len(r["bateu"]) == 32 and len(r["valor_diferente"]) == 2 and len(r["so_um_lado"]) == 4, nome


# ---------- regras de código (checklist da fase) ----------

def test_codigo_sem_log_de_conteudo_nem_comando_perigoso():
    proibido = re.compile(r"\bprint\(|logging|logger|\beval\(|\bexec\(|pickle|yaml\.load|shell=True|os\.system|subprocess|\.save\(")
    for arquivo in sorted((RAIZ / "core").glob("*.py")):
        if arquivo.name == "security.py":
            continue
        achados = [(n, l.strip()) for n, l in enumerate(arquivo.read_text(encoding="utf-8").splitlines(), 1)
                   if proibido.search(l) and not l.strip().startswith("#")]
        # `wb.save(saida)` no modelo grava em memória (BytesIO), não em disco
        achados = [a for a in achados if not (arquivo.name == "modelo_excel.py" and "wb.save(saida)" in a[1])]
        assert not achados, (arquivo.name, achados)


# ---------- rotas da tela Conferir (Fase 4) ----------

def _cliente():
    from app import app
    return app.test_client()


def _envio(lancamentos, extrato):
    return {"lancamentos": (io.BytesIO(lancamentos), "qualquer-nome.xlsx"), "extrato": (io.BytesIO(extrato), "outro-nome.ofx")}


def test_rota_processar_com_amostras_devolve_os_3_grupos():
    r = _cliente().post("/conferir/processar", data=_envio((AMOSTRAS / "lancamentos_omie_exemplo.xlsx").read_bytes(),
                        (AMOSTRAS / "extrato_bancoA_exemplo.ofx").read_bytes()), content_type="multipart/form-data")
    j = r.get_json()
    assert r.status_code == 200 and r.headers["Cache-Control"] == "no-store"
    assert (len(j["bateu"]), len(j["valor_diferente"]), len(j["so_um_lado"])) == (32, 2, 4)


def test_rota_exemplo_devolve_os_3_grupos():
    j = _cliente().post("/conferir/exemplo").get_json()
    assert j["resumo"]["bateram"] == 52 and len(j["valor_diferente"]) == 2 and len(j["so_um_lado"]) == 5


def test_rota_processar_erros_amigaveis():
    c = _cliente()
    ofx = (AMOSTRAS / "extrato_bancoA_exemplo.ofx").read_bytes()
    r = c.post("/conferir/processar", data={"extrato": (io.BytesIO(ofx), "a.ofx")}, content_type="multipart/form-data")
    assert r.status_code == 400 and r.get_json()["campo"] == "lancamentos"
    r = c.post("/conferir/processar", data=_envio(b"texto solto\n", ofx), content_type="multipart/form-data")
    j = r.get_json()
    assert r.status_code == 422 and j["campo"] == "lancamentos" and "Traceback" not in j["erro"]
    r = c.post("/conferir/processar", data=_envio(b"x" * (5 * 1024 * 1024 + 1), ofx), content_type="multipart/form-data")
    assert r.status_code == 413 and "5 MB" in r.get_json()["erro"]
    r = c.post("/conferir/processar", data=_envio(b"x" * (11 * 1024 * 1024), ofx), content_type="multipart/form-data")
    assert r.status_code == 413 and r.is_json


def test_upload_nao_grava_nem_usa_nome_do_arquivo_e_dom_so_por_textcontent():
    app_py = (RAIZ / "app.py").read_text(encoding="utf-8")
    assert ".save(" not in app_py and ".filename" not in app_py and "render_template_string" not in app_py
    js = (RAIZ / "static" / "js" / "conferir.js").read_text(encoding="utf-8")
    assert not re.search(r"innerHTML|insertAdjacentHTML|outerHTML|document\.write|\balert\(|\bconfirm\(|\bprompt\(", js)


def test_como_funciona_responde_com_os_3_blocos_sem_travessao():
    r = _cliente().get("/como-funciona")
    html = r.get_data(as_text=True)
    assert r.status_code == 200 and html.count('class="cd"') == 3
    assert "nada fica guardado" in html.lower() and "—" not in html


if __name__ == "__main__":
    testes = [(n, f) for n, f in sorted(globals().items()) if n.startswith("test_") and callable(f)]
    falhas = 0
    for nome, funcao in testes:
        try:
            funcao()
            print(f"ok     {nome}")
        except Exception as e:  # noqa: BLE001 - o runner mostra qualquer falha e segue
            falhas += 1
            print(f"FALHOU {nome}: {type(e).__name__}: {e}")
    print(f"\n{len(testes) - falhas} de {len(testes)} testes passaram")
    sys.exit(1 if falhas else 0)
